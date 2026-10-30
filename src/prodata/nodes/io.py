"""Ferramentas de entrada e saída."""

from __future__ import annotations

import os
import time
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Literal

import polars as pl
from pydantic import BaseModel, Field

from prodata.connectors import sql
from prodata.connectors.queries import read_definition, source_query_id
from prodata.errors import InvalidInputError
from prodata.forms import ui
from prodata.ingestion.datasets import resolve_inside, sanitize_filename
from prodata.ingestion.options import ReadOptions
from prodata.ingestion.readers import read_table
from prodata.nodes.base import NodeContext, NodeOutput, Port, ToolConfig, tool

# ------------------------------------------------------------------ Input Data


class InputDataConfig(ReadOptions, ToolConfig):
    dataset: str = Field(
        "",
        title="Fonte",
        description="Arquivo ou consulta de banco cadastrados na aba Dados.",
        json_schema_extra=ui(widget="dataset"),
    )
    limit: int = Field(0, ge=0, title="Limite de linhas", description="0 = todas as linhas.")


def _source_signature(ctx: NodeContext, cfg: InputDataConfig) -> str:
    if not cfg.dataset:
        return ""
    query_id = source_query_id(cfg.dataset)
    if query_id is None:
        return ctx.datasets.signature(cfg.dataset)
    if ctx.queries is None:
        return ""
    query = ctx.queries.get(query_id)
    # Sem cache, um carimbo de tempo força nova consulta (o banco pode ter mudado).
    return f"q{query.digest()}" + ("" if query.cache_results else f"t{time.time_ns()}")


def _source_annotation(cfg: InputDataConfig) -> str:
    if not cfg.dataset:
        return "Escolha uma fonte"
    return "Consulta de banco" if source_query_id(cfg.dataset) else cfg.dataset


@tool(
    type="io.input_data",
    name="Input Data",
    category="inout",
    icon="database",
    description="Lê qualquer fonte da aba Dados: arquivos (CSV, Excel, Parquet, JSON, Arrow) ou consultas de banco.",
    config=InputDataConfig,
    inputs=[],
    outputs=[Port("out", "O", "Dados")],
    annotate=_source_annotation,
    fingerprint=_source_signature,
)
def input_data(ctx: NodeContext, cfg: InputDataConfig) -> NodeOutput:
    if not cfg.dataset:
        raise ValueError("Escolha uma fonte no campo 'Fonte'.")
    query_id = source_query_id(cfg.dataset)
    if query_id is not None:
        if ctx.queries is None or ctx.connections is None:
            raise ValueError("Consultas de banco não estão disponíveis neste ambiente.")
        query = ctx.queries.get(query_id)
        df = read_definition(query, ctx.connections, limit=cfg.limit or None)
        return NodeOutput(tables={"out": df}, extra={"source": query.name, "source_read_at": _now()})
    path = ctx.datasets.resolve(cfg.dataset)
    info = ctx.datasets.info_of(cfg.dataset)
    df = read_table(path, cfg, n_rows=cfg.limit or None)
    # Quando o arquivo foi modificado e se é vinculado: o Configurar mostra de quando são os dados (spec 017).
    extra = {
        "source_modified": info.modified_at.isoformat() if info.modified_at else None,
        "source_read_at": _now(),
    }
    if info.linked_path:
        extra["source_linked"] = info.linked_path
    return NodeOutput(tables={"out": df}, extra=extra)


def _now() -> str:
    return datetime.now(UTC).isoformat(timespec="seconds")


# ------------------------------------------------------------------ Text Input


class TextTable(BaseModel):
    columns: list[str] = Field(default_factory=lambda: ["coluna_1"])
    rows: list[list[Any]] = Field(default_factory=lambda: [[""]])


class TextInputConfig(ToolConfig):
    table: TextTable = Field(default_factory=TextTable, title="Tabela", json_schema_extra=ui(widget="grid"))


def _infer_column(values: list[Any]) -> pl.Series:
    """Converte uma coluna digitada em texto para inteiro, decimal, booleano ou texto."""
    cleaned = [None if v is None or str(v).strip() == "" else str(v).strip() for v in values]
    present = [v for v in cleaned if v is not None]
    raw = pl.Series(cleaned, dtype=pl.String)
    if not present:
        return raw
    if all(v.lower() in ("true", "false", "verdadeiro", "falso", "sim", "não", "nao") for v in present):
        truthy = raw.str.to_lowercase().is_in(["true", "verdadeiro", "sim"])
        return pl.select(pl.when(raw.is_null()).then(None).otherwise(truthy)).to_series()
    as_int = raw.cast(pl.Int64, strict=False)
    if as_int.null_count() == raw.null_count():
        return as_int
    normalized = raw.str.replace_all(r"^(-?\d+),(\d+)$", "$1.$2")
    as_float = normalized.cast(pl.Float64, strict=False)
    if as_float.null_count() == raw.null_count():
        return as_float
    return raw


@tool(
    type="io.text_input",
    name="Text Input",
    category="inout",
    icon="table",
    description="Tabela pequena digitada ou colada direto no fluxo (de-para, parâmetros, metas).",
    config=TextInputConfig,
    inputs=[],
    outputs=[Port("out", "O", "Tabela")],
    annotate=lambda c: f"{len(c.table.rows)} linha(s) × {len(c.table.columns)} coluna(s)",
)
def text_input(ctx: NodeContext, cfg: TextInputConfig) -> NodeOutput:
    columns = [c.strip() or f"coluna_{i + 1}" for i, c in enumerate(cfg.table.columns)]
    if len(set(columns)) != len(columns):
        raise ValueError("Os nomes de coluna da tabela precisam ser únicos.")
    width = len(columns)
    rows = [(list(r) + [None] * width)[:width] for r in cfg.table.rows]
    data = {name: _infer_column([r[i] for r in rows]).alias(name) for i, name in enumerate(columns)}
    return NodeOutput(tables={"out": pl.DataFrame(list(data.values()))})


# ------------------------------------------------------------------ Browse


class BrowseConfig(ToolConfig):
    notes: str = Field("", title="Anotações", json_schema_extra=ui(widget="textarea"))


@tool(
    type="io.browse",
    name="Browse",
    category="inout",
    icon="eye",
    description="Ponto de inspeção: repassa os dados sem alteração para conferir perfil e qualidade.",
    config=BrowseConfig,
)
def browse(ctx: NodeContext, cfg: BrowseConfig) -> NodeOutput:
    return NodeOutput(tables={"out": ctx.input()})


# ------------------------------------------------------------------ Output Data

ExportFormat = Literal["parquet", "csv", "excel", "json", "arrow"]
_EXTENSIONS: dict[str, str] = {
    "parquet": ".parquet",
    "csv": ".csv",
    "excel": ".xlsx",
    "json": ".json",
    "arrow": ".arrow",
}


class OutputDataConfig(ToolConfig):
    filename: str = Field("resultado", title="Nome do arquivo", description="A extensão segue o formato.")
    folder: str = Field(
        "",
        title="Pasta de destino",
        description="Vazio = pasta padrão de exportações do PRO.DATA.",
        json_schema_extra=ui(widget="folder"),
    )
    format: ExportFormat = Field(
        "excel",
        title="Formato",
        json_schema_extra=ui(
            labels={
                "parquet": "Parquet",
                "csv": "CSV (;)",
                "excel": "Excel",
                "json": "JSON",
                "arrow": "Arrow IPC",
            }
        ),
    )
    sheet_name: str = Field(
        "Dados",
        title="Nome da planilha",
        json_schema_extra=ui(show_if={"format": ["excel"]}),
    )
    if_exists: Literal["ask", "replace", "new_file", "add_sheet"] = Field(
        "ask",
        title="Se o arquivo já existir",
        description='"Adicionar planilha" só vale para Excel: vários Output Data podem montar um arquivo só.',
        json_schema_extra=ui(
            labels={
                "ask": "Perguntar ao executar",
                "replace": "Substituir o arquivo",
                "new_file": "Criar um novo arquivo (data e hora no nome)",
                "add_sheet": "Adicionar ou atualizar esta planilha (Excel)",
            },
        ),
    )


_BAD_SHEET_CHARS = set("[]:*?/\\")


def sheet_title(raw: str) -> str:
    """Nome de planilha aceito pelo Excel: 1 a 31 caracteres, sem colchetes, : * ? / ou barra invertida."""
    name = raw.strip() or "Dados"
    bad = sorted(set(name) & _BAD_SHEET_CHARS)
    if bad:
        raise InvalidInputError(f"O nome da planilha não pode ter {' '.join(bad)}.")
    if len(name) > 31:
        raise InvalidInputError("O nome da planilha tem no máximo 31 caracteres.")
    return name


def write_excel_sheet(df: pl.DataFrame, path: Path, sheet: str, add: bool) -> None:
    """Grava a planilha. Com ``add`` e arquivo existente, mantém as outras planilhas (e a posição desta)."""
    if not (add and path.exists()):
        df.write_excel(path, worksheet=sheet, autofit=True)
        return
    from openpyxl import load_workbook

    try:
        book = load_workbook(path)
    except Exception as exc:  # arquivo corrompido ou que não é xlsx: mensagem acionável
        raise InvalidInputError(
            f"Não foi possível abrir {path.name} para adicionar a planilha: {exc}"
        ) from exc
    position = len(book.sheetnames)
    if sheet in book.sheetnames:
        position = book.sheetnames.index(sheet)
        del book[sheet]
    ws = book.create_sheet(sheet, position)
    ws.append(df.columns)
    for row in df.iter_rows():
        ws.append(list(row))
    for i, col in enumerate(df.columns, start=1):  # largura aproximada, como o autofit do write_excel
        sample = [len(str(v)) for v in df[col].head(200).to_list() if v is not None]
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(60, max([len(col), *sample]) + 2)
    tmp = path.with_name(path.name + ".tmp")
    book.save(tmp)
    os.replace(tmp, path)


def export_path(exports_dir: Path, cfg: OutputDataConfig) -> Path:
    stem = Path(sanitize_filename(cfg.filename or "resultado")).stem
    name = stem + _EXTENSIONS[cfg.format]
    if not cfg.folder.strip():
        return resolve_inside(exports_dir, name)
    # Pasta escolhida pelo usuário (constituição IV, exceção da spec 009): absoluta e já existente.
    folder = Path(cfg.folder.strip())
    if not folder.is_absolute():
        raise InvalidInputError("A pasta de destino precisa ser um caminho completo, como D:\\Relatorios.")
    if not folder.is_dir():
        raise InvalidInputError(f"A pasta de destino não existe: {folder}")
    return folder.resolve() / name


def new_file_path(path: Path) -> Path:
    """Mesmo nome com data e hora (e um número, se ainda assim existir): o arquivo antigo fica intacto."""
    stamp = datetime.now().strftime("%Y-%m-%d %Hh%M")
    candidate = path.with_name(f"{path.stem} {stamp}{path.suffix}")
    n = 2
    while candidate.exists():
        candidate = path.with_name(f"{path.stem} {stamp} ({n}){path.suffix}")
        n += 1
    return candidate


@tool(
    type="io.output_data",
    name="Output Data",
    category="inout",
    icon="download",
    description="Grava o resultado em Parquet, CSV, Excel, JSON ou Arrow, na pasta padrão ou numa pasta escolhida.",
    config=OutputDataConfig,
    annotate=lambda c: (
        f"{c.filename}{_EXTENSIONS[c.format]}" + (f" · {c.sheet_name}" if c.format == "excel" else "")
    ),
    side_effect=True,
)
def output_data(ctx: NodeContext, cfg: OutputDataConfig) -> NodeOutput:
    df = ctx.input()
    if cfg.if_exists == "add_sheet" and cfg.format != "excel":
        raise InvalidInputError(
            '"Adicionar planilha" só vale para Excel. Escolha outra opção em "Se o arquivo já existir".'
        )
    path = export_path(ctx.exports_dir, cfg)
    existed = path.exists()
    # "Perguntar" sem ninguém para responder (ex.: execução pela API): nunca sobrescreve em silêncio (spec 017).
    if existed and cfg.if_exists in ("ask", "new_file"):
        path = new_file_path(path)
    action = (
        "criado"
        if not existed
        else {"replace": "substituído", "add_sheet": "planilha atualizada"}.get(cfg.if_exists, "novo arquivo")
    )
    match cfg.format:
        case "csv":
            df.write_csv(path, separator=";")
        case "excel":
            write_excel_sheet(df, path, sheet_title(cfg.sheet_name), cfg.if_exists == "add_sheet")
        case "json":
            df.write_json(path)
        case "arrow":
            df.write_ipc(path, compression="zstd")
        case _:
            df.write_parquet(path)
    extra = {"exported_file": path.name, "export_action": action}
    if cfg.folder.strip():
        extra["exported_path"] = str(path)
    return NodeOutput(tables={"out": df}, extra=extra)


# ------------------------------------------------------------------ Database Input


class DatabaseInputConfig(ToolConfig):
    connection: str = Field("", title="Conexão", json_schema_extra=ui(widget="connection"))
    mode: Literal["table", "query"] = Field(
        "table", title="Ler", json_schema_extra=ui(labels={"table": "Uma tabela", "query": "Consulta SQL"})
    )
    table: str = Field(
        "",
        title="Tabela",
        description="esquema.tabela",
        json_schema_extra=ui(widget="db_table", show_if={"mode": ["table"]}),
    )
    query: str = Field(
        "",
        title="Consulta",
        description="Somente SELECT/WITH: a conexão é somente leitura.",
        json_schema_extra=ui(widget="db_query", show_if={"mode": ["query"]}),
    )
    limit: int = Field(0, ge=0, title="Limite de linhas", description="0 = todas as linhas.")
    cache_results: bool = Field(
        False,
        title="Reaproveitar resultado",
        description="Desligado: o banco é consultado de novo a cada execução.",
        json_schema_extra=ui(group="Avançado"),
    )


def _database_freshness(ctx: NodeContext, cfg: DatabaseInputConfig) -> str:
    # Sem cache, um carimbo de tempo força nova consulta (o banco pode ter mudado).
    return "cached" if cfg.cache_results else f"t{time.time_ns()}"


def _database_annotation(cfg: DatabaseInputConfig) -> str:
    if cfg.mode == "table":
        return cfg.table or "Escolha uma tabela"
    return " ".join(cfg.query.split())[:50] or "Escreva a consulta"


@tool(
    type="io.database_input",
    name="Database Input",
    category="inout",
    icon="server",
    description="(Antiga) Lê uma tabela ou consulta SQL. Use o Input Data com uma consulta da aba Dados.",
    config=DatabaseInputConfig,
    inputs=[],
    outputs=[Port("out", "O", "Dados")],
    annotate=_database_annotation,
    fingerprint=_database_freshness,
    hidden=True,
)
def database_input(ctx: NodeContext, cfg: DatabaseInputConfig) -> NodeOutput:
    if ctx.connections is None:
        raise ValueError("Conexões com banco não estão disponíveis neste ambiente.")
    if not cfg.connection:
        raise ValueError("Escolha uma conexão.")
    conn = ctx.connections.get(cfg.connection)
    if cfg.mode == "table":
        if not cfg.table:
            raise ValueError("Escolha uma tabela.")
        query = sql.table_query(conn, cfg.table)
    else:
        query = cfg.query
    df = sql.read_sql(conn, ctx.connections.password(conn.id), query, limit=cfg.limit or None)
    return NodeOutput(tables={"out": df})
