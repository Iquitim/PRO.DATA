"""Spec 015 (RF-1505, RF-1507): Output Data em Excel, com nome da planilha e várias planilhas num arquivo."""

from pathlib import Path
from typing import Any

import polars as pl
import pytest

from prodata.errors import InvalidInputError
from prodata.ingestion import DatasetStore
from prodata.nodes import REGISTRY, NodeContext

CLIENTES = pl.DataFrame({"id": [1, 2], "nome": ["Ana", "Bia"]})
VENDAS = pl.DataFrame({"id": [1, 1, 2], "valor": [10.5, 20.0, 7.25]})


def export(tmp_path: Path, df: pl.DataFrame, **config: Any) -> Path:
    spec = REGISTRY["io.output_data"]
    exports = tmp_path / "exports"
    exports.mkdir(exist_ok=True)
    ctx = NodeContext({"in": df}, {}, DatasetStore(tmp_path / "d", 10), exports)
    spec.run(ctx, spec.config.model_validate({"filename": "relatorio", **config}))
    return exports / "relatorio.xlsx"


def sheets(path: Path) -> list[str]:
    from openpyxl import load_workbook

    return load_workbook(path, read_only=True).sheetnames


def test_excel_is_the_default_format() -> None:
    assert REGISTRY["io.output_data"].config().format == "excel"


def test_sheets_add_up_in_one_file(tmp_path: Path) -> None:
    path = export(tmp_path, CLIENTES, sheet_name="Clientes")
    export(tmp_path, VENDAS, sheet_name="Vendas", if_exists="add_sheet")
    assert sheets(path) == ["Clientes", "Vendas"]
    read = pl.read_excel(path, sheet_name="Vendas")
    assert read["valor"].to_list() == [10.5, 20.0, 7.25]
    assert pl.read_excel(path, sheet_name="Clientes")["nome"].to_list() == ["Ana", "Bia"]


def test_updating_a_sheet_keeps_its_position(tmp_path: Path) -> None:
    path = export(tmp_path, CLIENTES, sheet_name="Clientes")
    export(tmp_path, VENDAS, sheet_name="Vendas", if_exists="add_sheet")
    export(tmp_path, CLIENTES.head(1), sheet_name="Clientes", if_exists="add_sheet")
    assert sheets(path) == ["Clientes", "Vendas"]
    assert pl.read_excel(path, sheet_name="Clientes").height == 1


def test_replace_overwrites_the_whole_file(tmp_path: Path) -> None:
    export(tmp_path, CLIENTES, sheet_name="Clientes")
    path = export(tmp_path, VENDAS, sheet_name="Vendas")
    assert sheets(path) == ["Vendas"]


@pytest.mark.parametrize("bad", ["Vendas/2026", "a" * 32, "[x]"])
def test_invalid_sheet_names_are_explained(tmp_path: Path, bad: str) -> None:
    with pytest.raises((InvalidInputError, ValueError), match="planilha"):
        export(tmp_path, CLIENTES, sheet_name=bad)
